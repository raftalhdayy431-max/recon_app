import sys
from pathlib import Path
from decimal import Decimal

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.loader import auto_map, build_records, load_table
from app.engine import (Settings, reconcile, STATUS_MATCHED, STATUS_DIFF, STATUS_DUP,
                        STATUS_MISSING_A, STATUS_MISSING_B, STATUS_POSSIBLE)
from app.normalize import parse_amount, norm_ref, parse_date
from tools.make_sample import make


def _run(tmp_path, **kw):
    pa, pb = make(str(tmp_path))
    ta, tb = load_table(pa), load_table(pb)
    ma, mb = auto_map(ta), auto_map(tb)
    ra, _ = build_records(ta, ma, "A")
    rb, _ = build_records(tb, mb, "B")
    return reconcile(ra, rb, Settings(**kw)), ta, tb, ma, mb, ra, rb


def test_parse_helpers():
    assert parse_amount("١٢٣٫٥٠") == Decimal("123.50")
    assert parse_amount("(1,250.75)") == Decimal("-1250.75")
    assert parse_amount("1.234,56") == Decimal("1234.56")
    assert parse_amount("500-") == Decimal("-500")
    assert norm_ref(" 000123.0 ") == "123"
    assert norm_ref(70012345.0) == "70012345"
    assert str(parse_date("٢٠٢٦/٠٩/٠٥")) == "2026-09-05"


def test_header_detection_and_mapping(tmp_path):
    res, ta, tb, ma, mb, ra, rb = _run(tmp_path)
    assert ta.header_row == 4
    assert ma["ref"] == "رقم الحوالة" and ma["debit"] == "مدين" and ma["credit"] == "دائن"
    assert ma["date"] == "التاريخ" and ma["party"] == "اسم العميل" and ma["currency"] == "العملة"
    assert mb["ref"] == "Reference" and mb["amount"] == "Amount" and mb["fee"] == "Commission"
    assert mb["date"] == "Value Date" and mb["party"] == "Beneficiary"
    # صف الإجمالي لا يُحسب
    assert all("إجمالي" not in r.desc for r in ra)


def test_integrity_and_statuses(tmp_path):
    res, *_ = _run(tmp_path)
    assert all(v == 0 for v in res.integrity.values()), res.integrity
    c = res.counts
    assert c[STATUS_MISSING_B] >= 2
    assert c[STATUS_MISSING_A] >= 2
    assert c[STATUS_DUP] == 2
    assert c[STATUS_DIFF] >= 6
    assert c[STATUS_MATCHED] >= 30
    assert sum(c.values()) == len(res.rows)


def test_each_record_used_once(tmp_path):
    res, *_, ra, rb = _run(tmp_path)
    ua = [r.a.uid for r in res.rows if r.a]
    ub = [r.b.uid for r in res.rows if r.b]
    assert len(ua) == len(set(ua)) == len(ra)
    assert len(ub) == len(set(ub)) == len(rb)


def test_explanations_present(tmp_path):
    res, *_ = _run(tmp_path)
    text = " | ".join(r.reason_text for r in res.rows)
    assert "الرسوم" in text
    assert "تبديل أرقام" in text or "الفاصلة العشرية" in text
    assert "فرق التاريخ" in text
    assert "اختلاف العملة" in text
    assert "داخل البيان" in text or any(r.method == "المرجع داخل البيان" for r in res.rows)
    assert any(r.status == STATUS_POSSIBLE for r in res.rows)
    assert all(r.action for r in res.rows)


def test_sign_opposite_auto():
    from app.loader import Record
    mk = lambda side, row, ref, amt: Record(side, row, ref, ref, "", None, Decimal(amt), "SAR",
                                            Decimal(0), "", "")
    a = [mk("A", i, f"R{i}000", 100 * i) for i in range(1, 6)]
    b = [mk("B", i, f"R{i}000", -100 * i) for i in range(1, 6)]
    res = reconcile(a, b)
    assert res.rel == -1
    assert res.counts[STATUS_MATCHED] == 5
    assert all(v == 0 for v in res.integrity.values())


def test_export_roundtrip(tmp_path):
    from openpyxl import load_workbook
    from app.exporter import export_report
    pa, pb = make(str(tmp_path))
    ta, tb = load_table(pa), load_table(pb)
    ra, sa = build_records(ta, auto_map(ta), "A")
    rb, sb = build_records(tb, auto_map(tb), "B")
    res = reconcile(ra, rb)
    out = export_report(str(tmp_path / "r.xlsx"), res, ta, tb, sa + sb)
    wb = load_workbook(out)
    assert wb.sheetnames == ["الملخص", "جدول المطابقة", "جدول الفروقات", "للمراجعة",
                             "بيانات المحاسب", "بيانات الصرافة"]
    ws = wb["جدول المطابقة"]
    assert ws.max_row == 4 + len(res.rows)
    assert ws.auto_filter.ref
    # مجموع جدول الفروقات لكل عملة = الفرق الصافي
    wd = wb["جدول الفروقات"]
    tot = {}
    for row in wd.iter_rows(min_row=5, values_only=True):
        if row[0] is None:
            continue
        tot[row[3]] = tot.get(row[3], 0) + row[6]
    for cur, v in tot.items():
        assert abs(v - float(res.net_diff(cur))) < 0.05, (cur, v, res.net_diff(cur))


def test_manual_link(tmp_path):
    from app.engine import manual_link
    res, *_ = _run(tmp_path)
    ra = next(r for r in res.rows if r.status == STATUS_MISSING_B)
    rb = next(r for r in res.rows if r.status == STATUS_MISSING_A)
    n_before = len(res.rows)
    new = manual_link(res, ra, rb)
    assert len(res.rows) == n_before - 1
    assert new.status in (STATUS_MATCHED, STATUS_DIFF)
    assert res.counts[STATUS_MISSING_B] == 1 and res.counts[STATUS_MISSING_A] == 1
    assert all(v == 0 for v in res.integrity.values()), res.integrity
